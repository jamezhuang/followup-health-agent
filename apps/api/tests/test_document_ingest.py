import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from app.document_ingest import parse_document, prepare


class DocumentIngestTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)

        self.base = Path(self.directory.name)
        self.raw = self.base / "knowledge_raw"
        self.raw.mkdir()

        self.manifest = self.base / "sources.json"

    def write_manifest(self, **changes):
        entry = {
            "id": "TEST-01",
            "file": "test.txt",
            "scope": "public",
            "title": "入库测试",
            "institution": "测试机构",
            "url": "https://example.org/source",
            "published_at": None,
            "collected_at": "2026-10-06",
        }

        entry.update(changes)

        self.manifest.write_text(
            json.dumps([entry], ensure_ascii=False),
            encoding="utf-8",
        )

    def test_chinese_utf8_and_repeatable_output(self):
        (self.raw / "test.txt").write_text(
            "中文资料\n第二行",
            encoding="utf-8-sig",
        )

        self.write_manifest()

        first = prepare(self.manifest, self.raw)

        self.assertEqual(
            first,
            prepare(self.manifest, self.raw),
        )

        self.assertEqual(
            first[0][0]["text"],
            "中文资料\n第二行",
        )

    def test_private_scope_is_rejected(self):
        self.write_manifest(scope="private")

        with self.assertRaises(ValueError):
            prepare(self.manifest, self.raw)

    def test_path_outside_raw_dir_is_rejected(self):
        (self.base / "outside.txt").write_text(
            "outside",
            encoding="utf-8",
        )

        self.write_manifest(file="../outside.txt")

        with self.assertRaisesRegex(ValueError, "目录内"):
            prepare(self.manifest, self.raw)

    def test_empty_text_is_rejected(self):
        (self.raw / "test.txt").write_text(
            "  \n",
            encoding="utf-8",
        )

        self.write_manifest()

        with self.assertRaisesRegex(ValueError, "正文为空"):
            prepare(self.manifest, self.raw)

    def test_duplicate_ids_are_rejected(self):
        self.write_manifest()

        entries = json.loads(self.manifest.read_text(encoding="utf-8"))

        self.manifest.write_text(
            json.dumps(entries * 2),
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "不能重复"):
            prepare(self.manifest, self.raw)

    def test_docx_preserves_body_order_and_row_locations(self):
        document = Document()
        document.add_paragraph("开头")

        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "说明"

        document.add_paragraph("结尾")

        path = self.raw / "test.docx"
        document.save(path)

        unit = parse_document(path)[0]

        self.assertEqual(
            unit["text"],
            "开头\n\n字段 | 说明\n\n结尾",
        )

        block = unit["blocks"][1]

        self.assertEqual(
            (block["table"], block["row"]),
            (1, 1),
        )

        self.assertEqual(
            unit["text"][block["start"] : block["end"]],
            "字段 | 说明",
        )

    def test_empty_pdf_page_is_rejected(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)

        path = self.raw / "empty.pdf"

        with path.open("wb") as stream:
            writer.write(stream)

        with self.assertRaisesRegex(ValueError, "第 1 页"):
            parse_document(path)


if __name__ == "__main__":
    unittest.main()
